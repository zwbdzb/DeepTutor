"""Tests for deeptutor.utils.document_extractor."""

from __future__ import annotations

import base64
import io
import zipfile

from docx import Document as DocxDocument
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
import pytest

from deeptutor.utils import document_extractor as document_extractor_module
from deeptutor.utils.document_extractor import (
    MAX_DOC_BYTES,
    MAX_EXTRACTED_CHARS_PER_DOC,
    CorruptDocumentError,
    DocumentTooLargeError,
    EmptyDocumentError,
    UnsupportedDocumentError,
    extract_documents_from_records,
    extract_epub_spine,
    extract_text_from_bytes,
    extract_text_from_path,
    is_document_extension,
    normalize_epub_archive,
)

# ---------------------------------------------------------------------------
# Fixtures — generate office docs on the fly
# ---------------------------------------------------------------------------


def _make_docx(paragraphs: list[str]) -> bytes:
    doc = DocxDocument()
    for p in paragraphs:
        doc.add_paragraph(p)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _make_xlsx(sheets: dict[str, list[list[object]]]) -> bytes:
    wb = Workbook()
    default = wb.active
    first = True
    for name, rows in sheets.items():
        ws = default if first else wb.create_sheet()
        ws.title = name
        for row in rows:
            ws.append(row)
        first = False
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _make_pptx(slides_text: list[list[str]]) -> bytes:
    prs = Presentation()
    for slide_texts in slides_text:
        slide = prs.slides.add_slide(prs.slide_layouts[5])  # blank-ish layout
        for i, text in enumerate(slide_texts):
            tb = slide.shapes.add_textbox(Inches(1), Inches(1 + i * 0.5), Inches(6), Inches(0.5))
            tb.text_frame.text = text
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


_CONTAINER_XML = (
    '<?xml version="1.0"?>'
    '<container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">'
    '<rootfiles><rootfile full-path="{opf}" media-type="application/oebps-package+xml"/>'
    "</rootfiles></container>"
)


def _make_epub(
    chapters: dict[str, str],
    spine: list[str] | None = None,
    *,
    opf_dir: str = "OEBPS",
    with_container: bool = True,
    with_opf: bool = True,
    wrapper: str = "",
    with_macosx: bool = False,
) -> bytes:
    """Build a minimal EPUB in memory.

    ``chapters`` maps member names (relative to ``opf_dir``) to XHTML body
    markup. ``spine`` is an ordered subset of chapter keys controlling the
    reading order; it defaults to the dict order.

    ``wrapper`` nests the whole book under one directory and ``with_macosx``
    adds AppleDouble resource forks — together, what macOS Finder's "Compress"
    produces.
    """
    opf_path = f"{opf_dir}/content.opf"
    manifest = "".join(
        f'<item id="ch{i}" href="{name}" media-type="application/xhtml+xml"/>'
        for i, name in enumerate(chapters)
    )
    ids = {name: f"ch{i}" for i, name in enumerate(chapters)}
    spine_ids = spine if spine is not None else list(chapters)
    spine_xml = "".join(f'<itemref idref="{ids[name]}"/>' for name in spine_ids)
    opf = (
        '<?xml version="1.0"?>'
        '<package xmlns="http://www.idpf.org/2007/opf" version="3.0">'
        f"<manifest>{manifest}</manifest><spine>{spine_xml}</spine></package>"
    )
    root = f"{wrapper}/" if wrapper else ""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr(f"{root}mimetype", "application/epub+zip")
        if with_container:
            zf.writestr(f"{root}META-INF/container.xml", _CONTAINER_XML.format(opf=opf_path))
        if with_opf:
            zf.writestr(f"{root}{opf_path}", opf)
        for name, body in chapters.items():
            zf.writestr(
                f"{root}{opf_dir}/{name}",
                f'<html xmlns="http://www.w3.org/1999/xhtml"><body>{body}</body></html>',
            )
            if with_macosx:
                # AppleDouble: the content file's name and extension, binary
                # resource-fork bytes inside.
                zf.writestr(f"__MACOSX/{opf_dir}/._{name}", b"\x00\x05\x16\x07" + b"\x00" * 60)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# is_document_extension
# ---------------------------------------------------------------------------


class TestIsDocumentExtension:
    def test_office(self) -> None:
        assert is_document_extension("foo.pdf")
        assert is_document_extension("foo.DOCX")
        assert is_document_extension("report.xlsx")
        assert is_document_extension("deck.pptx")
        assert is_document_extension("book.epub")

    def test_text_and_code(self) -> None:
        # Any extension in FileTypeRouter.TEXT_EXTENSIONS should be supported.
        assert is_document_extension("notes.txt")
        assert is_document_extension("readme.md")
        assert is_document_extension("module.py")
        assert is_document_extension("config.yaml")
        assert is_document_extension("data.json")
        assert is_document_extension("index.html")
        assert is_document_extension("table.csv")

    def test_unsupported(self) -> None:
        assert not is_document_extension("foo.png")
        assert not is_document_extension("foo.zip")
        assert not is_document_extension("foo.exe")
        assert not is_document_extension("foo")
        assert not is_document_extension("")


# ---------------------------------------------------------------------------
# extract_text_from_bytes — happy paths
# ---------------------------------------------------------------------------


class TestExtractDocx:
    def test_basic_paragraphs(self) -> None:
        data = _make_docx(["Hello world", "Second paragraph", ""])
        text = extract_text_from_bytes("doc.docx", data)
        assert "Hello world" in text
        assert "Second paragraph" in text

    def test_path_helper_can_disable_chat_truncation(self, tmp_path) -> None:
        data = _make_docx(["a" * (MAX_EXTRACTED_CHARS_PER_DOC + 10)])
        path = tmp_path / "long.docx"
        path.write_bytes(data)

        text = extract_text_from_path(path, max_chars=None)

        assert len(text) > MAX_EXTRACTED_CHARS_PER_DOC
        assert "truncated" not in text

    def test_ooxml_fallback_without_python_docx(self, monkeypatch: pytest.MonkeyPatch) -> None:
        monkeypatch.setattr(document_extractor_module, "DocxDocument", None)
        data = _make_docx(["Fallback paragraph", "第二段"])

        text = extract_text_from_bytes("doc.docx", data)

        assert "Fallback paragraph" in text
        assert "第二段" in text


class TestExtractXlsx:
    def test_multiple_sheets(self) -> None:
        data = _make_xlsx(
            {
                "Alpha": [["a1", "b1"], ["a2", 42]],
                "Beta": [["x", "y"]],
            }
        )
        text = extract_text_from_bytes("book.xlsx", data)
        assert "--- Sheet: Alpha ---" in text
        assert "--- Sheet: Beta ---" in text
        assert "a1" in text and "42" in text
        assert "x" in text

    def test_ooxml_fallback_without_openpyxl(self, monkeypatch: pytest.MonkeyPatch) -> None:
        monkeypatch.setattr(document_extractor_module, "load_workbook", None)
        data = _make_xlsx({"Alpha": [["name", "score"], ["alice", 98]]})

        text = extract_text_from_bytes("book.xlsx", data)

        assert "--- Sheet: Alpha ---" in text
        assert "alice" in text
        assert "98" in text


class TestExtractPptx:
    def test_basic_slides(self) -> None:
        data = _make_pptx([["Slide 1 title", "Slide 1 body"], ["Slide 2 only text"]])
        text = extract_text_from_bytes("deck.pptx", data)
        assert "--- Slide 1 ---" in text
        assert "--- Slide 2 ---" in text
        assert "Slide 1 title" in text
        assert "Slide 2 only text" in text

    def test_ooxml_fallback_without_python_pptx(self, monkeypatch: pytest.MonkeyPatch) -> None:
        monkeypatch.setattr(document_extractor_module, "PptxPresentation", None)
        data = _make_pptx([["Fallback slide", "第二行"]])

        text = extract_text_from_bytes("deck.pptx", data)

        assert "--- Slide 1 ---" in text
        assert "Fallback slide" in text
        assert "第二行" in text


class TestExtractEpub:
    def test_follows_spine_reading_order(self) -> None:
        data = _make_epub(
            {
                "chap1.xhtml": "<h1>Chapter One</h1><p>Alpha text.</p>",
                "chap2.xhtml": "<h1>Chapter Two</h1><p>Beta text.</p>",
            },
            spine=["chap2.xhtml", "chap1.xhtml"],
        )

        text = extract_text_from_bytes("book.epub", data)

        assert text.index("Chapter Two") < text.index("Chapter One")
        assert "Alpha text." in text
        assert "Beta text." in text

    def test_inline_markup_keeps_word_boundaries(self) -> None:
        data = _make_epub({"c.xhtml": "<p>Hello <b>world</b>. <i>Nice</i> day.</p>"})

        text = extract_text_from_bytes("book.epub", data)

        assert "Hello world. Nice day." in text

    def test_scripts_and_styles_are_dropped(self) -> None:
        data = _make_epub(
            {
                "c.xhtml": (
                    "<style>body{color:red}</style><script>var x=1;</script><p>Visible only.</p>"
                )
            }
        )

        text = extract_text_from_bytes("book.epub", data)

        assert "Visible only." in text
        assert "color:red" not in text
        assert "var x" not in text

    def test_falls_back_to_archive_order_without_container(self) -> None:
        data = _make_epub(
            {"a.xhtml": "<p>First member.</p>", "b.xhtml": "<p>Second member.</p>"},
            with_container=False,
        )

        text = extract_text_from_bytes("book.epub", data)

        assert "First member." in text
        assert "Second member." in text

    def test_falls_back_to_archive_order_without_opf(self) -> None:
        data = _make_epub({"only.xhtml": "<p>Solo chapter.</p>"}, with_opf=False)

        text = extract_text_from_bytes("book.epub", data)

        assert "Solo chapter." in text

    def test_a_finder_compressed_epub_still_reads_as_the_book_it_is(self) -> None:
        """macOS "Compress" wraps the book in a folder and adds ``__MACOSX``.

        Both together used to defeat the package lookup: the container was no
        longer at the archive root, so resolution fell back to matching file
        extensions, which picked up the AppleDouble forks as chapters. The
        reader then reported "Could not load this section" on binary members
        that were never part of the book (#1447).
        """
        data = _make_epub(
            {
                "index_split_000.xhtml": "<h1>Chapter One</h1><p>Alpha text.</p>",
                "index_split_001.xhtml": "<h1>Chapter Two</h1><p>Beta text.</p>",
            },
            spine=["index_split_001.xhtml", "index_split_000.xhtml"],
            wrapper="MyBook",
            with_macosx=True,
        )

        units, _ = extract_epub_spine(data, "book.epub")

        assert [unit.title for unit in units] == ["Chapter Two", "Chapter One"]
        assert all("__MACOSX" not in unit.href for unit in units)

    def test_resource_forks_are_not_chapters_even_without_a_package(self) -> None:
        """The extension-matching fallback must not read AppleDouble bytes."""
        data = _make_epub(
            {"a.xhtml": "<p>First member.</p>"},
            with_container=False,
            with_macosx=True,
        )

        units, _ = extract_epub_spine(data, "book.epub")

        assert [unit.href for unit in units] == ["OEBPS/a.xhtml"]

    def test_archive_normalization_preserves_a_root_level_book(self) -> None:
        data = _make_epub({"a.xhtml": "<p>Readable.</p>"})

        assert normalize_epub_archive(data, "book.epub") is data

    def test_archive_normalization_rejects_a_bad_zip(self) -> None:
        with pytest.raises(CorruptDocumentError, match="failed to open"):
            normalize_epub_archive(b"not a zip", "book.epub")

    def test_archive_normalization_repairs_a_finder_package(self) -> None:
        data = _make_epub(
            {"a.xhtml": "<p>Readable.</p>"},
            wrapper="MyBook",
            with_macosx=True,
        )

        normalized = normalize_epub_archive(data, "book.epub")

        with zipfile.ZipFile(io.BytesIO(normalized)) as zf:
            infos = zf.infolist()
            mimetype = zf.read("mimetype")

        assert [info.filename for info in infos] == [
            "mimetype",
            "META-INF/container.xml",
            "OEBPS/content.opf",
            "OEBPS/a.xhtml",
        ]
        assert infos[0].compress_type == zipfile.ZIP_STORED
        assert mimetype == b"application/epub+zip"

    def test_malformed_xhtml_uses_tolerant_html_fallback(self) -> None:
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as zf:
            zf.writestr("mimetype", "application/epub+zip")
            zf.writestr("META-INF/container.xml", _CONTAINER_XML.format(opf="OEBPS/content.opf"))
            zf.writestr(
                "OEBPS/content.opf",
                '<?xml version="1.0"?>'
                '<package xmlns="http://www.idpf.org/2007/opf" version="3.0">'
                '<manifest><item id="c0" href="bad.xhtml"/>'
                '<item id="c1" href="good.xhtml"/></manifest>'
                '<spine><itemref idref="c0"/><itemref idref="c1"/></spine></package>',
            )
            zf.writestr("OEBPS/bad.xhtml", "<html><body><p>Broken &nbsp; chapter</p></body></html>")
            zf.writestr("OEBPS/good.xhtml", "<html><body><p>Readable chapter.</p></body></html>")

        text = extract_text_from_bytes("book.epub", buf.getvalue())

        assert "Readable chapter." in text
        assert "Broken chapter" in text

    def test_rejects_suspicious_compression_ratio(self) -> None:
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", compression=zipfile.ZIP_DEFLATED) as zf:
            zf.writestr("mimetype", "application/epub+zip")
            zf.writestr("chapter.xhtml", f"<p>{'A' * 1_000_000}</p>")

        with pytest.raises(DocumentTooLargeError, match="compression ratio"):
            extract_text_from_bytes("book.epub", buf.getvalue())

    def test_rejects_excessive_member_count(self, monkeypatch: pytest.MonkeyPatch) -> None:
        monkeypatch.setattr(document_extractor_module, "_EPUB_MAX_MEMBERS", 2)
        data = _make_epub({"chapter.xhtml": "<p>text</p>"})

        with pytest.raises(DocumentTooLargeError, match="too many archive members"):
            extract_text_from_bytes("book.epub", data)

    def test_bad_header_raises_corrupt(self) -> None:
        with pytest.raises(CorruptDocumentError):
            extract_text_from_bytes("book.epub", b"not a zip at all")

    def test_zip_without_text_raises_empty(self) -> None:
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as zf:
            zf.writestr("mimetype", "application/epub+zip")
        with pytest.raises(EmptyDocumentError):
            extract_text_from_bytes("book.epub", buf.getvalue())


class TestExtractTextLike:
    def test_plain_txt(self) -> None:
        text = extract_text_from_bytes("note.txt", "hello world\nline two".encode("utf-8"))
        assert "hello world" in text
        assert "line two" in text

    def test_python_source(self) -> None:
        src = b"def greet(name: str) -> str:\n    return f'hi {name}'\n"
        text = extract_text_from_bytes("greet.py", src)
        assert "def greet" in text

    def test_json(self) -> None:
        text = extract_text_from_bytes("data.json", b'{"x": 42, "y": "ok"}')
        assert '"x": 42' in text

    def test_csv(self) -> None:
        text = extract_text_from_bytes("table.csv", b"a,b,c\n1,2,3\n")
        assert "a,b,c" in text
        assert "1,2,3" in text

    def test_markdown(self) -> None:
        text = extract_text_from_bytes("doc.md", b"# Heading\n\nBody.\n")
        assert "# Heading" in text

    def test_utf8_with_bom(self) -> None:
        # The candidate chain has utf-8 before utf-8-sig (same order as KB
        # pipeline), so BOM-prefixed bytes decode as utf-8 and the BOM is
        # retained. Mirror that behavior here.
        text = extract_text_from_bytes("note.txt", b"\xef\xbb\xbfhello")
        assert "hello" in text

    def test_gbk_fallback(self) -> None:
        # "你好" in GBK
        text = extract_text_from_bytes("note.txt", "你好".encode("gbk"))
        assert text == "你好"

    def test_svg(self) -> None:
        svg = (
            b'<?xml version="1.0"?>'
            b'<svg xmlns="http://www.w3.org/2000/svg" width="100" height="100">'
            b'<circle cx="50" cy="50" r="40" fill="red"/>'
            b'<text x="50" y="55">Hello</text>'
            b"</svg>"
        )
        text = extract_text_from_bytes("logo.svg", svg)
        assert "<svg" in text
        assert "<circle" in text
        assert "Hello" in text


class TestExtractPdf:
    def test_minimal_pdf(self) -> None:
        # Build a minimal valid PDF via pymupdf (dependency already in project).
        pytest.importorskip("fitz")
        import fitz  # noqa: WPS433

        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((72, 72), "Hello PDF world")
        buf = io.BytesIO()
        doc.save(buf)
        doc.close()
        data = buf.getvalue()

        text = extract_text_from_bytes("sample.pdf", data)
        assert "Hello PDF world" in text
        assert "--- Page 1 ---" in text


# ---------------------------------------------------------------------------
# Failure modes
# ---------------------------------------------------------------------------


class TestFailureModes:
    def test_unsupported_extension(self) -> None:
        with pytest.raises(UnsupportedDocumentError):
            extract_text_from_bytes("foo.zip", b"\x00\x00")

    def test_empty_bytes(self) -> None:
        with pytest.raises(EmptyDocumentError):
            extract_text_from_bytes("foo.docx", b"")

    def test_too_large(self) -> None:
        fake = b"PK\x03\x04" + b"\x00" * (MAX_DOC_BYTES + 1)
        with pytest.raises(DocumentTooLargeError):
            extract_text_from_bytes("foo.docx", fake)

    def test_pdf_magic_mismatch(self) -> None:
        with pytest.raises(CorruptDocumentError):
            extract_text_from_bytes("foo.pdf", b"this is not a pdf")

    def test_ooxml_magic_mismatch(self) -> None:
        with pytest.raises(CorruptDocumentError):
            extract_text_from_bytes("foo.docx", b"not an office file")

    def test_corrupt_docx(self) -> None:
        # OOXML header but garbage body
        with pytest.raises(CorruptDocumentError):
            extract_text_from_bytes("foo.docx", b"PK\x03\x04" + b"\x00" * 512)

    def test_empty_docx_no_text(self) -> None:
        data = _make_docx([])  # no paragraphs
        with pytest.raises(EmptyDocumentError):
            extract_text_from_bytes("foo.docx", data)


# ---------------------------------------------------------------------------
# Truncation
# ---------------------------------------------------------------------------


class TestTruncation:
    def test_long_docx_is_truncated(self) -> None:
        # single paragraph of 250k chars → well over the 200k per-doc cap
        long_text = "a" * 250_000
        data = _make_docx([long_text])
        text = extract_text_from_bytes("big.docx", data)
        assert len(text) <= MAX_EXTRACTED_CHARS_PER_DOC + 200  # allow notice suffix
        assert "truncated" in text


# ---------------------------------------------------------------------------
# extract_documents_from_records
# ---------------------------------------------------------------------------


class TestExtractDocumentsFromRecords:
    def test_mixed_image_and_doc(self) -> None:
        docx_bytes = _make_docx(["hello there"])
        docx_b64 = base64.b64encode(docx_bytes).decode()
        image_b64 = base64.b64encode(b"\x89PNG\r\n\x1a\n").decode()

        records = [
            {
                "type": "image",
                "filename": "pic.png",
                "base64": image_b64,
                "mime_type": "image/png",
                "url": "",
            },
            {
                "type": "file",
                "filename": "note.docx",
                "base64": docx_b64,
                "mime_type": "",
                "url": "",
            },
        ]

        doc_texts, updated = extract_documents_from_records(records)

        assert len(doc_texts) == 1
        assert "[File: note.docx]" in doc_texts[0]
        assert "hello there" in doc_texts[0]

        # image record untouched
        assert updated[0]["base64"] == image_b64
        # doc record base64 cleared, extracted_chars set
        assert updated[1]["base64"] == ""
        assert updated[1]["extracted_chars"] > 0

    def test_unsupported_record_is_passthrough(self) -> None:
        records = [
            {"type": "file", "filename": "foo.zip", "base64": "AAAA", "mime_type": "", "url": ""}
        ]
        doc_texts, updated = extract_documents_from_records(records)
        assert doc_texts == []
        assert updated[0]["base64"] == "AAAA"  # untouched — not a doc extension

    def test_failed_extraction_emits_error_marker(self) -> None:
        records = [
            {
                "type": "file",
                "filename": "bad.pdf",
                "base64": base64.b64encode(b"not a pdf").decode(),
                "mime_type": "",
                "url": "",
            }
        ]
        doc_texts, updated = extract_documents_from_records(records)
        assert len(doc_texts) == 1
        assert "bad.pdf" in doc_texts[0]
        assert "could not be read" in doc_texts[0]
        assert updated[0]["base64"] == ""  # stripped even on failure

    def test_invalid_base64_emits_error_marker(self) -> None:
        records = [
            {
                "type": "file",
                "filename": "bad.docx",
                "base64": "!!!not base64!!!",
                "mime_type": "",
                "url": "",
            }
        ]
        doc_texts, updated = extract_documents_from_records(records)
        # invalid base64 with validate=False may silently decode or emit error — both
        # paths end up as an error marker since resulting bytes won't pass magic check
        assert len(doc_texts) == 1
        assert "bad.docx" in doc_texts[0]

    def test_limits_come_from_settings_layer(self, monkeypatch) -> None:
        """extract_documents_from_records honors the configured policy."""
        from deeptutor.services.config import runtime_settings as rs

        # Tiny caps prove the configured values flow through (defaults would
        # accept everything here).
        def set_limits(max_file: int, max_total: int) -> None:
            monkeypatch.setattr(
                rs,
                "get_chat_attachment_limits",
                lambda: rs.ChatAttachmentLimits(
                    max_file_bytes=max_file,
                    max_total_bytes=max_total,
                    max_chars_per_doc=100_000,
                    max_chars_total=100_000,
                ),
            )

        def record(name: str, payload: bytes) -> dict:
            return {
                "type": "file",
                "filename": name,
                "base64": base64.b64encode(payload).decode(),
                "mime_type": "",
                "url": "",
            }

        # Per-file cap: 5 bytes, generous total.
        set_limits(max_file=5, max_total=1000)
        doc_texts, updated = extract_documents_from_records(
            [record("big.txt", b"0123456789"), record("ok.txt", b"hello")]
        )
        assert "per-file limit" in doc_texts[0]
        assert doc_texts[1] == "[File: ok.txt]\nhello"
        assert updated[1]["extracted_chars"] == 5

        # Per-turn total: 8 bytes — the second 5-byte file blows it.
        set_limits(max_file=1000, max_total=8)
        doc_texts, _ = extract_documents_from_records(
            [record("a.txt", b"hello"), record("b.txt", b"world")]
        )
        assert doc_texts[0] == "[File: a.txt]\nhello"
        assert "quota exceeded" in doc_texts[1]

    def test_char_budget_comes_from_settings_layer(self, monkeypatch) -> None:
        from deeptutor.services.config import runtime_settings as rs

        monkeypatch.setattr(
            rs,
            "get_chat_attachment_limits",
            lambda: rs.ChatAttachmentLimits(
                max_file_bytes=10 * 1024 * 1024,
                max_total_bytes=25 * 1024 * 1024,
                max_chars_per_doc=100_000,
                max_chars_total=10,
            ),
        )
        records = [
            {
                "type": "file",
                "filename": "long.txt",
                "base64": base64.b64encode(b"a" * 40).decode(),
                "mime_type": "",
                "url": "",
            }
        ]
        doc_texts, updated = extract_documents_from_records(records)
        assert "truncated" in doc_texts[0]
        assert updated[0]["extracted_chars"] <= 10 + len(
            "... (truncated, 40 chars total; turn quota hit)"
        )
